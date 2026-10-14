"""Dalolatnomalar bo'yicha fuqarolar ro'yxatini yaratish.

fuqarolar.json  ->  PDF'larni fuqaro ismi bilan qayta nomlash,
                    fuqarolar_royxati.xlsx va fuqarolar_royxati.json yaratish.

Ishlatish:
    python build.py ["<PDF papkasi>"]   (ko'rsatilmasa, dastur/ ning yuqorisidagi papka)
"""
import json
import os
import re
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
MUAMMOLI_DIR = "Muammoli"

# PINFL birinchi raqami: jins va tug'ilgan asr
PINFL_PREFIX = {"1": ("Erkak", 1800), "2": ("Ayol", 1800),
                "3": ("Erkak", 1900), "4": ("Ayol", 1900),
                "5": ("Erkak", 2000), "6": ("Ayol", 2000)}

COLUMNS = [
    ("№", 5),
    ("F.I.Sh", 36),
    ("PINFL (JShShIR)", 17),
    ("Jinsi", 8),
    ("Tug'ilgan sana", 13),
    ("Hozirgi mahallasi (ro'yxatdagi)", 30),
    ("Yangi mahallasi", 30),
    ("Yangi yashash manzili", 48),
    ("Telefon", 15),
    ("Dalolatnoma sanasi", 12),
    ("Izoh", 50),
    ("Muammo", 28),
    ("PDF fayl", 40),
]


def pinfl_info(pinfl):
    """PINFL'ni tekshiradi: (to'g'rimi, jinsi, tug'ilgan sana, xato matni)."""
    if not pinfl:
        return False, "", "", ""
    if not re.fullmatch(r"\d{14}", pinfl):
        return False, "", "", "PINFL 14 ta raqam emas"
    if pinfl[0] not in PINFL_PREFIX:
        return False, "", "", "PINFL birinchi raqami noto'g'ri"
    gender, century = PINFL_PREFIX[pinfl[0]]
    dd, mm, yy = int(pinfl[1:3]), int(pinfl[3:5]), int(pinfl[5:7])
    try:
        born = date(century + yy, mm, dd)
    except ValueError:
        return False, gender, "", "PINFL ichidagi sana noto'g'ri"
    weights = [7, 3, 1] * 5
    check = sum(int(d) * w for d, w in zip(pinfl[:13], weights)) % 10
    if check != int(pinfl[13]):
        return False, gender, born.strftime("%d.%m.%Y"), "PINFL nazorat raqami mos emas"
    return True, gender, born.strftime("%d.%m.%Y"), ""


def safe_filename(name):
    return re.sub(r'[<>:"/\\|?*]', "", name).strip()


def unique_name(folder, base, used):
    candidate, n = f"{base}.pdf", 2
    while candidate in used or (os.path.exists(os.path.join(folder, candidate)) and candidate not in used):
        candidate = f"{base} ({n}).pdf"
        n += 1
    used.add(candidate)
    return candidate


def muammo(r):
    """Fuqaro nega muammoli ekanini qaytaradi ('' bo'lsa muammo yo'q)."""
    if not r["yangi_mahalla"]:
        return "Yangi mahallasi aniqlanmagan"
    return ""


def rename_pdfs(folder, records):
    """PDF'ni fuqaro ismi bilan nomlaydi; muammolilarni Muammoli papkasiga o'tkazadi."""
    used = set()
    os.makedirs(os.path.join(folder, MUAMMOLI_DIR), exist_ok=True)
    for r in records:
        new = unique_name(folder, safe_filename(r["fish"]), used) if not r.get("yangi_fayl") else r["yangi_fayl"]
        target_rel = os.path.join(MUAMMOLI_DIR, new) if muammo(r) else new
        target = os.path.join(folder, target_rel)
        # fayl qayerda bo'lsa ham topiladi: eski nomda, asosiy papkada yoki Muammoli'da
        sources = [os.path.join(folder, r["eski_fayl"]), os.path.join(folder, new),
                   os.path.join(folder, MUAMMOLI_DIR, new)]
        if os.path.exists(target):
            status = "joyida"
        else:
            src = next((s for s in sources if os.path.exists(s)), None)
            if src:
                os.rename(src, target)
                status = "ko'chirildi"
            else:
                status = "FAYL TOPILMADI"
        r["yangi_fayl"] = new
        r["pdf_joyi"] = target_rel.replace("\\", "/")
        print(f"  [{status}] {r['eski_fayl']}  ->  {r['pdf_joyi']}")


def enrich(records):
    for r in records:
        ok, gender, born, err = pinfl_info(r["pinfl"])
        r["jinsi"] = gender
        if born and not r["tugilgan_sana"]:
            r["tugilgan_sana"] = born
        elif born and r["tugilgan_sana"] != born:
            err = (err + "; " if err else "") + f"PINFL'dagi sana ({born}) hujjatdagidan farq qiladi"
        r["pinfl_tekshiruv"] = "to'g'ri" if ok else (err or "PINFL yo'q")
        r["muammo"] = muammo(r)
        if err:
            r["izoh"] = (r["izoh"] + "; " if r["izoh"] else "") + err


def write_excel(path, records):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fuqarolar"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")

    for col, (title, width) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=col, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 32

    for i, r in enumerate(records, 1):
        row = [i, r["fish"], r["pinfl"], r["jinsi"], r["tugilgan_sana"], r["hozirgi_mahalla"],
               r["yangi_mahalla"], r["yangi_manzil"], r["telefon"], r["dalolatnoma_sanasi"],
               r["izoh"], r["muammo"], r["pdf_joyi"]]
        for col, value in enumerate(row, 1):
            c = ws.cell(row=i + 1, column=col, value=value)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = border
            if col == 3:
                c.number_format = "@"  # PINFL matn sifatida, 3.08E+13 bo'lib qolmasin
            if r["pinfl_tekshiruv"] != "to'g'ri" or r["muammo"]:
                c.fill = warn_fill

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(HERE)
    with open(os.path.join(HERE, "fuqarolar.json"), encoding="utf-8") as f:
        records = json.load(f)

    print("PDF fayllar qayta nomlanmoqda:")
    rename_pdfs(folder, records)
    enrich(records)

    # yangi_fayl manbaga ham yoziladi, shunda skriptni qayta ishga tushirish xavfsiz
    with open(os.path.join(HERE, "fuqarolar.json"), "w", encoding="utf-8") as f:
        source = [{k: r[k] for k in r if k not in ("jinsi", "pinfl_tekshiruv", "muammo", "pdf_joyi")} for r in records]
        json.dump(source, f, ensure_ascii=False, indent=2)

    xlsx = os.path.join(folder, "fuqarolar_royxati.xlsx")
    out_json = os.path.join(folder, "fuqarolar_royxati.json")
    write_excel(xlsx, records)
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

    problems = [r for r in records if r["muammo"]]
    mdir = os.path.join(folder, MUAMMOLI_DIR)
    write_excel(os.path.join(mdir, "muammoli_royxati.xlsx"), problems)
    with open(os.path.join(mdir, "muammoli_royxati.json"), "w", encoding="utf-8") as f:
        json.dump(problems, f, ensure_ascii=False, indent=2)
    print(f"\nMuammoli ({MUAMMOLI_DIR}/ papkasida): {len(problems)} ta")
    for r in problems:
        print(f"  - {r['fish']}: {r['muammo']}")

    bad = [r for r in records if r["pinfl_tekshiruv"] != "to'g'ri"]
    print(f"\nJami: {len(records)} ta fuqaro. PINFL muammoli/yo'q: {len(bad)} ta")
    for r in bad:
        print(f"  - {r['fish']}: {r['pinfl_tekshiruv']}")
    print(f"\nExcel: {xlsx}\nJSON:  {out_json}")


if __name__ == "__main__":
    main()
